"""
Convierte el Excel de formulación PAC (Anexo 6) al JSON de carga masiva.
Uso: python excel_to_pac_json.py MATRIZ_FORMULACION.xlsx output.json [--pac-id UUID]

Requiere: pip install openpyxl
"""
import openpyxl
import json
import sys
import uuid
from decimal import Decimal


class DecimalEncoder(json.JSONEncoder):
    def default(self, o):
        if isinstance(o, Decimal):
            return float(o)
        return super().default(o)


def resolve_merged_value(ws, row, col):
    """Devuelve el valor de una celda, resolviendo merged ranges."""
    cell = ws.cell(row=row, column=col)
    if cell.value is not None:
        return cell.value
    for mr in ws.merged_cells.ranges:
        if cell.coordinate in mr:
            return ws.cell(row=mr.min_row, column=mr.min_col).value
    return None


def safe_decimal(val):
    if val is None:
        return None
    try:
        return Decimal(str(val))
    except Exception:
        return None


def convert(excel_path, output_path, pac_id=None):
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # Buscar la hoja que contenga 'FormulacionPAC' o 'Anexo 6'
    target_sheet = None
    for name in wb.sheetnames:
        if 'formulacionpac' in name.lower() or 'anexo 6' in name.lower():
            target_sheet = name
            break
    if target_sheet is None:
        target_sheet = wb.sheetnames[0]
        print(f"AVISO: No se encontró hoja con 'FormulacionPAC'. Usando: {target_sheet}")

    ws = wb[target_sheet]

    # Mapeo de columnas esperado (ajustar según el Excel real):
    # A=línea estratégica, B=peso línea, C=código programa, D=nombre programa,
    # E=peso programa, F=código proyecto, G=nombre proyecto, H=peso proyecto,
    # I=vigencia, J=subdivisión, K=objetivo, L=peso objetivo,
    # M=código actividad, N=nombre actividad, O=peso actividad,
    # P=indicador, Q=tipo indicador, R=unidad de medida,
    # S-V=metas anuales (2024-2027), W-Z=recursos anuales (2024-2027)

    # Año base configurable
    YEARS = [2024, 2025, 2026, 2027]
    START_ROW = 3  # Primera fila de datos (ajustar según encabezados)

    COL = {
        'line_name': 1,
        'line_weight': 2,
        'prog_code': 3,
        'prog_name': 4,
        'prog_weight': 5,
        'proj_code': 6,
        'proj_name': 7,
        'proj_weight': 8,
        'vigencia': 9,
        'subdivision': 10,
        'objective': 11,
        'obj_weight': 12,
        'act_code': 13,
        'act_name': 14,
        'act_weight': 15,
        'indicator': 16,
        'indicator_type': 17,
        'unit': 18,
        'goal_start': 19,   # S=meta 2024, T=meta 2025, ...
        'budget_start': 23, # W=recurso 2024, X=recurso 2025, ...
    }

    lines = {}  # name -> BulkLineRequest dict
    current_line = None
    current_prog = None
    current_proj = None

    for row in range(START_ROW, ws.max_row + 1):
        line_name = resolve_merged_value(ws, row, COL['line_name'])
        prog_code = resolve_merged_value(ws, row, COL['prog_code'])
        proj_code = resolve_merged_value(ws, row, COL['proj_code'])
        act_code = resolve_merged_value(ws, row, COL['act_code'])

        # Propagar valores hacia abajo (celdas merged)
        if line_name:
            line_name = str(line_name).strip()
            if line_name and line_name not in lines:
                lines[line_name] = {
                    'name': line_name,
                    'weight': float(safe_decimal(resolve_merged_value(ws, row, COL['line_weight'])) or 0),
                    'programs': {}
                }
            if line_name:
                current_line = lines[line_name]

        if current_line is None:
            continue

        if prog_code:
            prog_code = str(prog_code).strip()
            prog_name = str(resolve_merged_value(ws, row, COL['prog_name']) or '').strip()
            if prog_code not in current_line['programs']:
                current_line['programs'][prog_code] = {
                    'code': prog_code,
                    'name': prog_name,
                    'weight': float(safe_decimal(resolve_merged_value(ws, row, COL['prog_weight'])) or 0),
                    'projects': {}
                }
            current_prog = current_line['programs'][prog_code]

        if current_prog is None:
            continue

        if proj_code:
            proj_code = str(proj_code).strip()
            proj_name = str(resolve_merged_value(ws, row, COL['proj_name']) or '').strip()
            if proj_code not in current_prog['projects']:
                current_prog['projects'][proj_code] = {
                    'code': proj_code,
                    'name': proj_name,
                    'weight': float(safe_decimal(resolve_merged_value(ws, row, COL['proj_weight'])) or 0),
                    'vigencia': int(resolve_merged_value(ws, row, COL['vigencia']) or 2026),
                    'subdivision': str(resolve_merged_value(ws, row, COL['subdivision']) or '').strip() or None,
                    'objective': str(resolve_merged_value(ws, row, COL['objective']) or '').strip() or None,
                    'objectiveWeight': float(safe_decimal(resolve_merged_value(ws, row, COL['obj_weight'])) or 0),
                    'activities': []
                }
            current_proj = current_prog['projects'][proj_code]

        if current_proj is None:
            continue

        if act_code:
            act_code = str(act_code).strip()
            act_name = str(resolve_merged_value(ws, row, COL['act_name']) or '').strip()

            goals = []
            budgets = []
            for i, year in enumerate(YEARS):
                goal_val = safe_decimal(ws.cell(row=row, column=COL['goal_start'] + i).value)
                budget_val = safe_decimal(ws.cell(row=row, column=COL['budget_start'] + i).value)
                if goal_val is not None:
                    goals.append({'year': year, 'value': float(goal_val)})
                if budget_val is not None:
                    budgets.append({'year': year, 'value': float(budget_val)})

            current_proj['activities'].append({
                'code': act_code,
                'name': act_name,
                'weight': float(safe_decimal(resolve_merged_value(ws, row, COL['act_weight'])) or 0),
                'measurementUnit': str(resolve_merged_value(ws, row, COL['unit']) or '').strip() or None,
                'indicatorDescription': str(resolve_merged_value(ws, row, COL['indicator']) or '').strip() or None,
                'indicatorType': str(resolve_merged_value(ws, row, COL['indicator_type']) or '').strip() or None,
                'sdgIndicatorIds': [],
                'annualGoals': goals,
                'annualBudgets': budgets
            })

    # Convertir dicts intermedios a listas
    result = {
        'pacId': pac_id or str(uuid.uuid4()),
        'strategicLines': []
    }
    for line_data in lines.values():
        programs_list = []
        for prog_data in line_data['programs'].values():
            projects_list = list(prog_data['projects'].values())
            programs_list.append({
                'code': prog_data['code'],
                'name': prog_data['name'],
                'weight': prog_data['weight'],
                'projects': projects_list
            })
        result['strategicLines'].append({
            'name': line_data['name'],
            'weight': line_data['weight'],
            'programs': programs_list
        })

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2, cls=DecimalEncoder)

    # Resumen
    total_lines = len(result['strategicLines'])
    total_progs = sum(len(l['programs']) for l in result['strategicLines'])
    total_projs = sum(len(p['projects']) for l in result['strategicLines'] for p in l['programs'])
    total_acts = sum(
        len(pj['activities'])
        for l in result['strategicLines']
        for p in l['programs']
        for pj in p['projects']
    )
    print(f"Conversión completada:")
    print(f"  Líneas estratégicas: {total_lines}")
    print(f"  Programas: {total_progs}")
    print(f"  Proyectos: {total_projs}")
    print(f"  Actividades: {total_acts}")
    print(f"  Archivo: {output_path}")


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Uso: python excel_to_pac_json.py MATRIZ.xlsx output.json [--pac-id UUID]")
        sys.exit(1)

    excel_file = sys.argv[1]
    output_file = sys.argv[2]
    pac_uuid = None
    if '--pac-id' in sys.argv:
        idx = sys.argv.index('--pac-id')
        pac_uuid = sys.argv[idx + 1]

    convert(excel_file, output_file, pac_uuid)
